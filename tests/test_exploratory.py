"""Scientific input, derived KPI and complete-grid contracts for policy mixes."""
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
import openpyxl

from pv_batteries import exploratory_config as cfg
from pv_batteries.exploratory import validate_kpi_store, scale_kpis, score_policy_mixes
from pv_batteries.exploratory_kpis import pv_self_coverage


def synthetic_store():
    rows = []
    for pid, policy in enumerate(cfg.all_policy_combos()):
        for sid in range(cfg.N_SAMPLES):
            ratio = .1 + .5 * pid / 127 + .01 * sid / 100
            rows.append({'policy_id': pid, 'sample_id': sid, **policy,
                         **{k: 1. + pid + sid / 100 for k in cfg.KPI_ORDER},
                         'pv_self_coverage': ratio, 'pv_self_consumed_gwh': ratio,
                         'electricity_demand_gwh': 1., 'completed': 1.})
    return pd.DataFrame(rows)


class ExploratoryTests(unittest.TestCase):
    def test_inputs_match_the_workbook(self):
        workbook = openpyxl.load_workbook(cfg.ROOT / 'Input and Output Streamlit.xlsx', read_only=True, data_only=True)
        try:
            sheet = workbook['Exploratory analysis']
            policy = {sheet.cell(r, 4).value: [sheet.cell(r, 2).value, sheet.cell(r, 3).value] for r in range(2, 9)}
            uncertain = {sheet.cell(r, 4).value: (sheet.cell(r, 2).value, sheet.cell(r, 3).value) for r in range(13, 30)}
            self.assertEqual(policy, cfg.POLICY_GRID)
            self.assertEqual(uncertain, cfg.UNCERTAINTY_BOUNDS)
            self.assertEqual(sheet['H7'].value, cfg.N_SAMPLES)
            self.assertEqual(len([sheet.cell(r, 1).value for r in range(32, 41)]), len(cfg.KPI_ORDER))
        finally:
            workbook.close()

    def test_samples_are_uniform_stratified_and_reproducible(self):
        samples = cfg.build_uncertainty_samples()
        cfg.validate_samples(samples)
        self.assertEqual(len(samples), 100)
        for name, (low, high) in cfg.UNCERTAINTY_BOUNDS.items():
            bins = np.floor((samples[name].to_numpy() - low) / (high - low) * 100).astype(int)
            np.testing.assert_array_equal(np.sort(bins), np.arange(100))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'samples.parquet'
            samples.to_parquet(path, index=False)
            cfg.validate_samples(pd.read_parquet(path))
            self.assertTrue(cfg.load_or_create_samples(path).equals(samples))
            cfg.build_uncertainty_samples(50).to_parquet(path, index=False)
            with self.assertRaises(ValueError):
                cfg.load_or_create_samples(path)
        for invalid in (samples.iloc[::-1], samples.assign(sample_id=0), cfg.build_uncertainty_samples(seed=1)):
            with self.assertRaises(ValueError):
                cfg.validate_samples(invalid)

    def test_coverage_converts_units_and_sums_districts_before_dividing(self):
        ratio, numerator, denominator = pv_self_coverage(.75, [1e6, 2e6])
        self.assertEqual((ratio, numerator, denominator), (.25, .75, 3.))
        self.assertEqual(pv_self_coverage(0, [1e6])[0], 0.)
        self.assertEqual(pv_self_coverage(3, [1e6, 2e6])[0], 1.)
        for numerator, demand in [(1, [0, 0]), (1, [np.nan]), (np.nan, [1e6]),
                                  (-1, [1e6]), (1, [-1, 2e6]), (2, [1e6]), (1, []),
                                  ([1], [1e6]), (1, [[1e6]])]:
            with self.assertRaises(ValueError):
                pv_self_coverage(numerator, demand)

    def test_complete_grid_ratio_and_policy_inputs_are_required(self):
        frame = synthetic_store()
        validate_kpi_store(frame)
        cases = [frame.iloc[:-1], pd.concat([frame.iloc[:-1], frame.iloc[:1]], ignore_index=True),
                 frame.drop(columns='pv_self_coverage')]
        for column, value in [('pv_self_coverage', np.nan), ('pv_self_coverage', 1e-7),
                              ('electricity_demand_gwh', 0), ('completed', 0),
                              (cfg.POLICY_ORDER[0], .5)]:
            bad = frame.copy()
            bad.loc[0, column] = value
            cases.append(bad)
        for bad in cases:
            with self.assertRaises(ValueError):
                validate_kpi_store(bad)

    def test_new_kpi_drives_the_weighted_ranking(self):
        scaled = scale_kpis(synthetic_store())
        weights = {k: float(k == 'pv_self_coverage') for k in cfg.KPI_ORDER}
        ranked = score_policy_mixes(scaled, weights)
        self.assertEqual(ranked.iloc[0]['policy_id'], 127)
        self.assertEqual(ranked.iloc[-1]['policy_id'], 0)
        self.assertTrue((ranked['n_samples'] == 100).all())
        self.assertAlmostEqual(ranked.iloc[0]['score'], scaled.loc[scaled.policy_id == 127, 'pv_self_coverage_s'].mean())
        scaled.loc[0, 'pv_self_coverage_s'] = np.nan
        with self.assertRaises(ValueError):
            score_policy_mixes(scaled, weights)

    def test_full_batch_requests_all_nine_kpi_inputs(self):
        from sure_julia.batch import jobs
        with patch('sure_julia.batch.pd.read_parquet', return_value=cfg.build_uncertainty_samples()):
            planned = jobs()
        self.assertEqual(len(planned), 14571)
        self.assertEqual(sum(len(j['destinations']) for j in planned), 14579)
        exploratory = [j for j in planned if any(d['section'] == 'exploratory' for d in j['destinations'])]
        self.assertEqual(len(exploratory), 12800)
        for job in exploratory:
            self.assertIn(cfg.SC_NUMERATOR, job['outputs'])
            self.assertIn(cfg.SC_DENOMINATOR, job['outputs'])


if __name__ == '__main__':
    unittest.main()
